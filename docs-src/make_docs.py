"""Generate the project PPT and report (DOCX) for Bihar Heritage.
Run: python3 make_docs.py <screenshot_dir> <output_dir>
"""
import sys, os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from docx import Document
from docx.shared import Pt as DPt, Inches as DIn, RGBColor as DRGB
from docx.enum.text import WD_ALIGN_PARAGRAPH

SHOTS, OUT = sys.argv[1], sys.argv[2]
APP = os.path.dirname(os.path.abspath(__file__)) + '/..'
URL = 'https://sanjay289.github.io/bihar-heritage/'
RED, ORANGE, CREAM, DARK = RGBColor(0xC0, 0x39, 0x2B), RGBColor(0xE6, 0x7E, 0x22), RGBColor(0xFF, 0xF8, 0xEE), RGBColor(0x33, 0x33, 0x33)

# ------------------------------------------------------------------ PPT
p = Presentation()
p.slide_width, p.slide_height = Inches(13.333), Inches(7.5)
BLANK = p.slide_layouts[6]


def bg(s, color):
    f = s.background.fill
    f.solid()
    f.fore_color.rgb = color


def text(s, x, y, w, h, lines, size=20, color=DARK, bold=False, align=PP_ALIGN.LEFT):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines if isinstance(lines, list) else [lines]):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.text = line
        para.alignment = align
        para.space_after = Pt(8)
        for r in para.runs:
            r.font.size = Pt(size)
            r.font.color.rgb = color
            r.font.bold = bold
    return tb


def bar(s, title):
    bg(s, CREAM)
    r = s.shapes.add_shape(1, 0, 0, p.slide_width, Inches(1.1))
    r.fill.solid()
    r.fill.fore_color.rgb = RED
    r.line.fill.background()
    text(s, 0.5, 0.2, 12, 0.8, title, 32, RGBColor(255, 255, 255), True)


def slide(title, bullets, image=None):
    s = p.slides.add_slide(BLANK)
    bar(s, title)
    w = 7.3 if image else 12.3
    text(s, 0.6, 1.5, w, 5.5, ['• ' + b for b in bullets], 22)
    if image and os.path.exists(image):
        s.shapes.add_picture(image, Inches(8.4), Inches(1.35), height=Inches(5.9))
    return s


# 1. Title
s = p.slides.add_slide(BLANK)
bg(s, RED)
if os.path.exists(f'{APP}/icon-512.png'):
    s.shapes.add_picture(f'{APP}/icon-512.png', Inches(5.67), Inches(0.6), height=Inches(2))
text(s, 0.5, 2.8, 12.3, 1, 'Bihar Heritage | बिहार धरोहर', 48, RGBColor(255, 255, 255), True, PP_ALIGN.CENTER)
text(s, 0.5, 3.9, 12.3, 1, 'An educational website and app about the history, people and culture of Bihar', 24, CREAM, False, PP_ALIGN.CENTER)
text(s, 0.5, 5.2, 12.3, 1, ['Presented by: ____________________', 'Class / College: ____________________', URL], 20, CREAM, False, PP_ALIGN.CENTER)

slide('The Problem', [
    'Many students know little about the rich history of their own state.',
    'Textbooks give few pages to Bihar\'s places, heroes and culture.',
    'Most learning material is only in Hindi or English, not in local languages like Bhojpuri, Maithili or Magahi.',
    'Children in villages often have weak internet, so heavy apps don\'t work for them.'])

slide('Our Solution', [
    'A free, simple website that also works like a mobile app.',
    'Teaches Bihar\'s places, great people, culture and all 38 districts.',
    'Available in 5 languages: Hindi, English, Bhojpuri, Maithili and Magahi.',
    'A fun quiz with levels and scores makes learning a game.',
    'Works offline once opened, and can be installed on the home screen.'],
    f'{SHOTS}/s1.png')

slide('Features', [
    '🏛️ Places: Nalanda, Bodh Gaya, Vaishali, Rajgir and more (12)',
    '👑 People: Chanakya, Ashoka, Aryabhata, Dr. Rajendra Prasad and more (12)',
    '🎨 Culture: Madhubani art, Chhath, litti chokha, makhana and more (12)',
    '🗺️ Districts: all 38 districts in 9 divisions, with search',
    '❓ Quiz: 54 questions, Easy / Hard / Mixed levels',
    '📷 Real photos with proper credits'],
    f'{SHOTS}/s2.png')

slide('The Quiz', [
    '10 random questions every round, options shuffled each time.',
    'Right answer turns green, wrong turns red, phone vibrates.',
    'Progress bar and live score.',
    'Best score saved for each level.',
    'Questions cover places, people, culture and districts.'],
    f'{SHOTS}/s4.png')

slide('Technology Used', [
    'HTML, CSS and JavaScript (no framework, very light)',
    'Progressive Web App (PWA): manifest + service worker for offline use and installation',
    'GitHub Pages for free hosting',
    'Wikimedia Commons for freely licensed photos',
    'Built and tested on an Android phone using Termux',
    'Whole app is about 1.3 MB, so it loads fast on slow networks'])

slide('How It Helps', [
    'Students learn about Bihar in their own mother tongue.',
    'Teachers can use it in history and social science classes.',
    'Builds pride in Bihar\'s heritage among young people.',
    'Useful for competitive exam preparation (BPSC, SSC general knowledge).',
    'Free, with no ads and no login.'])

slide('Future Plans', [
    'Add Angika and Bajjika languages.',
    'Audio narration for children who can\'t read yet.',
    'More photos and a short video for each place.',
    'Let schools add stories from their own district.',
    'Leaderboard for school quiz competitions.'])

s = p.slides.add_slide(BLANK)
bg(s, RED)
text(s, 0.5, 2.2, 12.3, 1, 'Thank You!  धन्यवाद!', 54, RGBColor(255, 255, 255), True, PP_ALIGN.CENTER)
text(s, 0.5, 3.6, 12.3, 1, ['Try it yourself:', URL], 26, CREAM, False, PP_ALIGN.CENTER)
text(s, 0.5, 6.5, 12.3, 0.6, 'Photos: Wikimedia Commons contributors (see "Photo credits" in the app)', 14, CREAM, False, PP_ALIGN.CENTER)

p.save(f'{OUT}/Bihar-Heritage-Presentation.pptx')

# ------------------------------------------------------------------ Report
d = Document()
st = d.styles['Normal']
st.font.name = 'Calibri'
st.font.size = DPt(12)


def h(t, lvl=1):
    x = d.add_heading(t, lvl)
    for r in x.runs:
        r.font.color.rgb = DRGB(0xC0, 0x39, 0x2B)


def para(t):
    d.add_paragraph(t)


def bullets(items):
    for i in items:
        d.add_paragraph(i, style='List Bullet')


t = d.add_heading('Project Report', 0)
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
x = d.add_paragraph()
x.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = x.add_run('Bihar Heritage | बिहार धरोहर\nAn Educational Website and App about Bihar')
r.bold = True
r.font.size = DPt(18)
x = d.add_paragraph()
x.alignment = WD_ALIGN_PARAGRAPH.CENTER
x.add_run('\nSubmitted by: ______________________\nClass / Roll No.: ______________________\n'
          'School / College: ______________________\nGuided by: ______________________\n\nLive link: ' + URL)
d.add_page_break()

h('Certificate')
para('This is to certify that the project "Bihar Heritage" has been completed by ______________________ '
     'of class ______ under my guidance during the session ______.')
para('\n\nSignature of Guide: ____________________          Signature of Principal: ____________________')
h('Acknowledgement')
para('I would like to thank my teacher ______________________ for guidance and support. I also thank my parents '
     'and friends who tested the app and gave suggestions, and the Wikimedia Commons contributors whose free photos '
     'are used in this project.')
d.add_page_break()

h('1. Introduction')
para('Bihar is one of the oldest centres of learning and culture in the world. It is the land of Nalanda and '
     'Vikramshila universities, of Buddha and Mahavira, of Chanakya, Ashoka and Aryabhata. Yet many students know '
     'very little about this heritage. "Bihar Heritage" is a free website and mobile app that teaches the history, '
     'great people, culture and districts of Bihar in a simple and fun way, in five languages.')

h('2. Objectives')
bullets([
    'To make learning about Bihar\'s heritage easy and interesting for students.',
    'To provide content in local languages: Bhojpuri, Maithili and Magahi, along with Hindi and English.',
    'To use a quiz so that students test and remember what they read.',
    'To build an app that works on cheap phones and without internet.',
    'To keep it completely free, with no ads and no login.'])

h('3. Problem Statement')
para('Students in Bihar mostly study general history, and local heritage gets little space. Many children are more '
     'comfortable in their mother tongue than in Hindi or English. Internet in villages is often slow or missing. '
     'So there is a need for a light, multilingual, offline learning tool about Bihar.')

h('4. Features of the App')
bullets([
    'Places: 12 important places such as Nalanda, Bodh Gaya, Vaishali, Rajgir, Golghar and Barabar Caves.',
    'People: 12 great personalities such as Chanakya, Ashoka, Aryabhata, Guru Gobind Singh and Dr. Rajendra Prasad.',
    'Culture: 12 topics such as Madhubani and Manjusha art, Chhath Puja, Sonepur Mela, litti chokha and makhana.',
    'Districts: all 38 districts of Bihar grouped into 9 divisions, each with what it is famous for.',
    'Quiz: 54 multiple-choice questions in Easy, Hard and Mixed levels, with score and best score.',
    'Languages: Hindi, English, Bhojpuri, Maithili and Magahi.',
    'Search: find any topic or district quickly.',
    'Real photos from Wikimedia Commons, with a credits page.',
    'Offline use and "Add to Home screen" installation (Progressive Web App).'])
for img in ('s1.png', 's2.png', 's4.png'):
    if os.path.exists(f'{SHOTS}/{img}'):
        d.add_picture(f'{SHOTS}/{img}', width=DIn(2.6))
        d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
para('Figure: Screenshots of the Places screen, the Districts screen (in Maithili) and the Quiz.')

h('5. Technology Used')
table = d.add_table(rows=1, cols=2)
table.style = 'Light Grid Accent 2'
table.rows[0].cells[0].text, table.rows[0].cells[1].text = 'Part', 'Technology'
for a, b in [('Structure', 'HTML5'), ('Design', 'CSS3 (mobile-first, responsive)'), ('Logic', 'JavaScript (no framework)'),
             ('Offline and install', 'Progressive Web App: Web App Manifest and Service Worker'),
             ('Storage', 'Browser localStorage (language choice and best scores)'),
             ('Hosting', 'GitHub Pages (free)'), ('Photos', 'Wikimedia Commons (free licences)'),
             ('Development', 'Android phone with Termux')]:
    c = table.add_row().cells
    c[0].text, c[1].text = a, b

h('6. How It Works')
bullets([
    'All content is stored in JavaScript files (data.js and data2.js), so no server or database is needed.',
    'When the user picks a language, every screen is redrawn in that language. If a line is not translated, Hindi is shown.',
    'The quiz picks 10 random questions from the chosen level and shuffles the options.',
    'The service worker saves all files on the phone the first time, so the app opens later even without internet.'])

h('7. Testing')
para('The app was tested in the Chromium browser with an automatic test that checked: all 5 language buttons, all '
     'cards in every language, photos loading, 38 districts in 9 divisions, division filter, search, and full quiz '
     'rounds in every level and several languages. All checks passed. It was also opened on an Android phone, both '
     'from the live link and from a saved file.')

h('8. Benefits')
bullets([
    'Students learn local history in their own language.',
    'Teachers can use it in class on a phone or projector.',
    'Helpful for general knowledge in exams like BPSC and SSC.',
    'Creates pride in Bihar\'s heritage.'])

h('9. Limitations')
bullets([
    'Bhojpuri, Maithili and Magahi translations should be checked by native speakers.',
    'Quiz questions show Hindi text in Bhojpuri, Maithili and Magahi modes (except the first set of Bhojpuri questions).',
    'There is no audio yet for children who cannot read.'])

h('10. Future Scope')
bullets([
    'Add Angika and Bajjika languages and audio narration.',
    'Add videos and more photos for each place and district.',
    'Let schools submit stories from their own district.',
    'Add a leaderboard for school quiz competitions.'])

h('11. Conclusion')
para('"Bihar Heritage" shows that a small, free and light app can help students learn about their state in their '
     'own language. It works on simple phones, even offline, and makes learning fun through a quiz. With more content '
     'and audio it can become a useful tool for schools across Bihar.')

h('12. References')
bullets([
    'Wikipedia articles on Bihar, its districts, places and personalities.',
    'Wikimedia Commons (photos, individual credits listed in the app).',
    'Government of Bihar official website (list of districts and divisions).',
    'MDN Web Docs (Progressive Web Apps).'])

d.save(f'{OUT}/Bihar-Heritage-Project-Report.docx')
print('ok')
